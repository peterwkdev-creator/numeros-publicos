"""Leitura da planilha de IDEB por município do INEP.

Módulo **puro de leitura**: recebe bytes ou caminho, devolve dados. A rede fica
em `baixar()`, com o transporte injetado, e o banco em `armazem.py`.

## As seis armadilhas, e onde cada uma é tratada

Nenhuma estoura. Cada uma publica número errado e fica quieta. As quatro
primeiras foram medidas na planilha antes de escrever código; **as duas últimas
só apareceram quando o leitor rodou contra o arquivo real** — a fixture, escrita
por quem já entendera as quatro, concordava por construção.

1. **`CO_MUNICIPIO` não é chave — `REDE` é a outra metade.** Cada município
   aparece de 1 a 4 vezes, e `Pública` **é o agregado** das outras, não uma rede
   a mais. Ver `REDES` e `REDE_AGREGADA`.
2. **Ausência é a string `"-"` ou `"ND"`** — 229.837 e 3.236 ocorrências. Ver
   `AUSENTE` e `_valor_e_marca()`.
3. **O cabeçalho está na linha 10.** As seis primeiras são timbre do
   Ministério. Ver `LINHA_CABECALHO`.
4. **Sem `User-Agent`, o INEP devolve HTTP 200 com corpo vazio.** Ver
   `baixar()` e `conferir_md5()`.
5. **Número vem como texto com ressalva colada:** `"206,95**"`, e `"ND*"` é
   ausência **com** ressalva. A legenda está no **rodapé** da planilha, não no
   cabeçalho. Ver `MARCAS`.
6. **2005, 2023 e 2025 não têm meta.** 2005 é a linha de base (as metas
   começam em 2007) e o primeiro ciclo do IDEB encerrou em 2021 — a Portaria
   26/2024 criou grupo para definir as novas. Exibir meta para 2023 ou 2025
   seria inventar uma que não existe; o leitor devolve `None` e quem apresenta
   tem de dizer por quê.
"""

from __future__ import annotations

import hashlib
import io
import re
import zipfile
from dataclasses import dataclass
from typing import Callable, Iterator

import openpyxl

BASE = "https://download.inep.gov.br/ideb/resultados"

# Sem isto o INEP responde 200 com zero byte -- nao 403, nao 404. O antidoto e
# duplo: mandar o cabecalho E conferir o MD5 que vem dentro do zip.
CABECALHOS = {"User-Agent": "Mozilla/5.0 (compatible; numerospublicos/1.0)"}

# As seis primeiras linhas sao timbre do Ministerio; a 7 e a 9 sao rotulos para
# humano ("Sigla da UF", "1o ao 5o ano"). A 10 tem os nomes de maquina.
LINHA_CABECALHO = 10

ETAPAS = {"anos_iniciais": "anos_iniciais", "anos_finais": "anos_finais"}

# `Publica` NAO e uma rede: e a soma das outras tres. Somar as linhas de um
# municipio contaria o mesmo aluno duas vezes.
REDE_AGREGADA = "Pública"
REDES = ("Municipal", "Estadual", "Federal", REDE_AGREGADA)

# A rede que este projeto le. O gasto com educacao exibido ao lado vem do
# orcamento DO MUNICIPIO; parear com um indice que inclui escola estadual
# atribuiria a prefeitura um resultado que ela nao administra nem financia.
REDE_PADRAO = "Municipal"

# Como a fonte escreve "nao ha valor". Nao e vazio, nao e zero, nao e nulo.
AUSENTE = {"-", "", "--", "ND"}

# As marcas de ressalva que o INEP cola no proprio numero -- `206,95**`. A
# legenda esta no RODAPE da planilha (linhas 14513-14514 da divulgacao de
# 2023), e nao no cabecalho, que e onde ninguem procura.
#
# Elas nao sao decoracao: sao aviso de qualidade que a fonte esta dando.
# Descartar o valor perderia um numero real; tirar o asterisco em silencio
# descartaria o aviso. Guardamos os dois.
MARCAS = {
    "*": "Média calculada a partir dos resultados das escolas nas avaliações "
         "estaduais, por extravio das provas retornadas.",
    "**": "Média calculada com participação inferior a 50% no SAEB, por "
          "extravio de provas.",
    "***": "Ver a nota informativa do INEP sobre a edição.",
}


class ArquivoVazio(RuntimeError):
    """HTTP 200 com corpo vazio -- o modo de falha do INEP sem `User-Agent`."""


class Md5NaoConfere(RuntimeError):
    """O zip traz o hash de cada planilha; se nao bate, o download corrompeu."""


@dataclass(frozen=True)
class LinhaIdeb:
    """Uma linha da planilha, já com as colunas de ano desdobradas em edições."""

    codigo_ibge: int
    uf: str
    nome: str
    rede: str
    etapa: str
    edicao: int
    observado: float | None
    projecao: float | None
    nota_portugues: float | None
    nota_matematica: float | None
    indicador_rendimento: float | None
    # As ressalvas que a fonte colou nas notas do Saeb, quando houve. `None`
    # quando o número veio limpo. Ver `MARCAS`.
    marca_portugues: str | None = None
    marca_matematica: str | None = None

    @property
    def ressalvas(self) -> tuple[str, ...]:
        """As ressalvas desta linha, já em português, sem repetir."""
        vistas = []
        for m in (self.marca_portugues, self.marca_matematica):
            texto = MARCAS.get(m or "")
            if texto and texto not in vistas:
                vistas.append(texto)
        return tuple(vistas)

    @property
    def vazia(self) -> bool:
        """Edição em que a rede não teve nenhum indicador calculado.

        Vale gravar mesmo assim? Não: linha só com ausência não é observação, é
        ruído — e 21.545 delas inflariam o banco sem dizer nada que o próprio
        silêncio já não diga.
        """
        return all(v is None for v in (
            self.observado, self.projecao, self.nota_portugues,
            self.nota_matematica, self.indicador_rendimento))


def url_de(etapa: str, ano: int = 2025) -> str:
    """A URL do zip de uma etapa.

    Só os anos mais recentes existem neste padrão: `..._2019.zip` e
    `..._2021.zip` dão **404**. Não é problema — o arquivo de 2025 traz as onze
    edições desde 2005 em colunas, então a série sai de um download só. Medido
    em 11/10/2026: as edições de 2005 a 2023 do arquivo de 2025 são idênticas
    às do de 2023 (observado e meta, linha a linha).
    """
    if etapa not in ETAPAS:
        raise ValueError(f"etapa desconhecida: {etapa!r}; use {sorted(ETAPAS)}")
    return f"{BASE}/divulgacao_{etapa}_municipios_{ano}.zip"


# O arquivo tem 23 MB e a conexao do INEP cai no meio dele. Medido em
# 03/09/2026: `WinError 10054` (conexao resetada pelo host remoto) na primeira
# tentativa, e as tres seguintes passando sem tocar em nada. Nao e o
# `User-Agent` -- e transporte mesmo, e vai repetir.
TENTATIVAS = 4
ESPERA_INICIAL = 2.0


def baixar(url: str, transporte: Callable[[str, dict], bytes],
           dormir: Callable[[float], None] | None = None,
           tentativas: int = TENTATIVAS) -> bytes:
    """Baixa o zip, com retentativa. Transporte e relógio injetados.

    **Corpo vazio é erro, não sucesso.** É exatamente o que o INEP devolve, com
    HTTP 200, quando falta o `User-Agent` — e um ingestor que aceite isso grava
    arquivo vazio e relata que deu tudo certo.

    **Corpo vazio NÃO é retentado.** Ele não é falha de rede: é a resposta que
    o servidor dá quando o pedido está errado, e repeti-lo produziria quatro
    respostas vazias idênticas e uma mensagem de erro pior.
    """
    import time

    dormir = dormir or time.sleep
    espera = ESPERA_INICIAL
    ultima: Exception | None = None
    for tentativa in range(1, tentativas + 1):
        try:
            dados = transporte(url, CABECALHOS)
        except Exception as e:   # noqa: BLE001 -- qualquer falha de transporte
            ultima = e
            if tentativa == tentativas:
                break
            dormir(espera)
            espera *= 2
            continue
        if not dados:
            raise ArquivoVazio(
                f"{url} devolveu 0 byte. O INEP responde 200 com corpo vazio "
                f"quando falta o cabeçalho User-Agent — conferir CABECALHOS."
            )
        return dados
    raise RuntimeError(
        f"{url} falhou em {tentativas} tentativas; a última foi "
        f"{type(ultima).__name__}: {ultima}"
    ) from ultima


def conferir_md5(zip_bytes: bytes) -> str:
    """Confere a planilha contra o MD5 que o próprio INEP embarca no zip.

    É a régua que a fonte oferece de graça, como o total declarado do RREO no
    painel fiscal. Devolve o nome do `.xlsx` conferido.

    O `md5.txt` tem uma linha por arquivo, no formato `<hash> *<nome>`.
    """
    z = zipfile.ZipFile(io.BytesIO(zip_bytes))
    nomes = z.namelist()
    arquivos_md5 = [n for n in nomes if n.endswith(".txt") and "md5" in n.lower()]
    if not arquivos_md5:
        raise Md5NaoConfere(
            f"o zip não trouxe md5.txt; arquivos presentes: {nomes}")

    esperados: dict[str, str] = {}
    for linha in z.read(arquivos_md5[0]).decode("utf-8", "replace").splitlines():
        m = re.match(r"\s*([0-9a-fA-F]{32})\s+\*?(.+?)\s*$", linha)
        if m:
            esperados[m.group(2)] = m.group(1).lower()

    xlsx = [n for n in nomes if n.endswith(".xlsx")]
    if not xlsx:
        raise Md5NaoConfere(f"o zip não trouxe .xlsx; presentes: {nomes}")
    alvo = xlsx[0]
    curto = alvo.rsplit("/", 1)[-1]
    if curto not in esperados:
        raise Md5NaoConfere(
            f"{curto} não aparece no md5.txt, que lista {sorted(esperados)}")

    calculado = hashlib.md5(z.read(alvo)).hexdigest()
    if calculado != esperados[curto]:
        raise Md5NaoConfere(
            f"{curto}: md5 calculado {calculado} != declarado "
            f"{esperados[curto]} — download corrompido ou truncado")
    return alvo


def extrair_xlsx(zip_bytes: bytes) -> bytes:
    """Os bytes da planilha, **depois** de conferir o MD5."""
    nome = conferir_md5(zip_bytes)
    return zipfile.ZipFile(io.BytesIO(zip_bytes)).read(nome)


def _valor_e_marca(v) -> tuple[float | None, str | None]:
    """Separa o número da ressalva que a fonte colou nele.

    Três coisas acontecem aqui, e cada uma existe por um motivo medido no
    arquivo real de 2023:

    1. **`"-"` e `"ND"` viram ausência, nunca zero.** 229.837 e 3.236
       ocorrências. Zero é um resultado; "não disponível" não é, e a diferença
       some para sempre se colapsarem.
    2. **Número vem como texto com vírgula decimal** (`"231,42"`), 23.990 vezes.
    3. **Asterisco é ressalva da fonte, não sujeira.** `"206,95**"` significa que
       a média foi calculada com participação inferior a 50% no SAEB. O valor é
       real e a ressalva também — devolvemos os dois. `"ND*"` é ausência **com**
       ressalva, e a ressalva sobrevive à ausência.
    """
    if v is None:
        return None, None
    if isinstance(v, (int, float)):
        return float(v), None

    texto = str(v).strip()
    marca = None
    corpo = texto
    while corpo.endswith("*"):
        corpo = corpo[:-1]
        marca = (marca or "") + "*"
    corpo = corpo.strip()

    if corpo in AUSENTE:
        return None, marca
    try:
        return float(corpo.replace(",", ".")), marca
    except ValueError as e:
        raise ValueError(
            f"valor não numérico e não reconhecido como ausência: {v!r}. "
            f"Se a fonte passou a usar outra marca, acrescentar em AUSENTE "
            f"(ausência) ou em MARCAS (ressalva), conferindo a legenda no "
            f"RODAPÉ da planilha."
        ) from e


def _numero(v) -> float | None:
    """Só o número, descartando a ressalva. Para quem não precisa dela."""
    return _valor_e_marca(v)[0]


def _cabecalho(ws) -> list[str]:
    """Os nomes de máquina, da linha 10 — não da primeira.

    Falha alto se a planilha mudar de forma: é o canário do sistema. O INEP
    reformata entre edições, e ler o timbre do Ministério como cabeçalho não
    estoura, só produz lixo.
    """
    for i, linha in enumerate(ws.iter_rows(values_only=True), start=1):
        if i == LINHA_CABECALHO:
            nomes = [("" if c is None else str(c).strip()) for c in linha]
            if not nomes or nomes[0] != "SG_UF":
                raise ValueError(
                    f"a linha {LINHA_CABECALHO} deveria começar com SG_UF e "
                    f"começa com {nomes[:1]!r}. A planilha do INEP mudou de "
                    f"forma — conferir antes de publicar qualquer número."
                )
            return nomes
    raise ValueError(f"a planilha tem menos de {LINHA_CABECALHO} linhas")


def ler(xlsx_bytes: bytes, etapa: str,
        redes: tuple[str, ...] = (REDE_PADRAO,)) -> Iterator[LinhaIdeb]:
    """Percorre a planilha e devolve uma `LinhaIdeb` por município/rede/edição.

    `redes` filtra na leitura. O padrão é só `Municipal`, e **`Pública` nunca
    deve ser lida junto com as outras**: ela é o agregado, e somá-las contaria
    o mesmo aluno duas vezes.
    """
    desconhecidas = set(redes) - set(REDES)
    if desconhecidas:
        raise ValueError(f"rede desconhecida: {sorted(desconhecidas)}")
    if REDE_AGREGADA in redes and len(redes) > 1:
        raise ValueError(
            f"{REDE_AGREGADA!r} é o AGREGADO das outras redes, não uma a mais. "
            f"Lê-la junto com {sorted(set(redes) - {REDE_AGREGADA})} contaria o "
            f"mesmo aluno duas vezes."
        )

    wb = openpyxl.load_workbook(io.BytesIO(xlsx_bytes), read_only=True,
                                data_only=True)
    ws = wb[wb.sheetnames[0]]
    nomes = _cabecalho(ws)
    col = {nome: i for i, nome in enumerate(nomes) if nome}

    # As edições saem dos próprios nomes de coluna, não de uma lista fixa: uma
    # divulgação nova acrescenta VL_OBSERVADO_2025 e o leitor a encontra sozinho.
    edicoes = sorted(
        int(m.group(1))
        for nome in nomes
        if (m := re.fullmatch(r"VL_OBSERVADO_(\d{4})", nome))
    )
    if not edicoes:
        raise ValueError(
            "nenhuma coluna VL_OBSERVADO_<ano> na linha de cabeçalho — a "
            f"planilha traz {nomes[:8]}...")

    def pega(linha, prefixo: str, ano: int) -> float | None:
        i = col.get(f"{prefixo}_{ano}")
        return None if i is None else _numero(linha[i])

    def pega_com_marca(linha, prefixo: str, ano: int):
        i = col.get(f"{prefixo}_{ano}")
        return (None, None) if i is None else _valor_e_marca(linha[i])

    escolhidas = set(redes)
    it = ws.iter_rows(min_row=LINHA_CABECALHO + 1, values_only=True)
    for linha in it:
        if not linha or linha[col["CO_MUNICIPIO"]] in (None, ""):
            continue
        rede = str(linha[col["REDE"]]).strip()
        if rede not in escolhidas:
            continue
        codigo = int(linha[col["CO_MUNICIPIO"]])
        uf = str(linha[col["SG_UF"]]).strip()
        nome = str(linha[col["NO_MUNICIPIO"]]).strip()
        for ano in edicoes:
            port, marca_port = pega_com_marca(linha, "VL_NOTA_PORTUGUES", ano)
            mat, marca_mat = pega_com_marca(linha, "VL_NOTA_MATEMATICA", ano)
            item = LinhaIdeb(
                codigo_ibge=codigo, uf=uf, nome=nome, rede=rede, etapa=etapa,
                edicao=ano,
                observado=pega(linha, "VL_OBSERVADO", ano),
                projecao=pega(linha, "VL_PROJECAO", ano),
                nota_portugues=port, nota_matematica=mat,
                indicador_rendimento=pega(linha, "VL_INDICADOR_REND", ano),
                marca_portugues=marca_port, marca_matematica=marca_mat,
            )
            if not item.vazia:
                yield item
    wb.close()
